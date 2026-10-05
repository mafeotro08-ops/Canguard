# -*- coding: utf-8 -*-
from fastapi import FastAPI, Request, Form, UploadFile, File
from typing import Optional
from fastapi.responses import HTMLResponse, RedirectResponse, JSONResponse
from fastapi.templating import Jinja2Templates
from fastapi.staticfiles import StaticFiles
from starlette.middleware.sessions import SessionMiddleware
import os
import re
import shutil
import secrets
import hashlib
from email.utils import parseaddr

try:
    from dotenv import load_dotenv
    load_dotenv(os.path.join(os.path.dirname(__file__), '..', '.env'))
except ImportError:
    pass

from datetime import date, timedelta
from .models import (
    verify_user, get_all_entidades, get_entidad_by_id,
    create_entidad, update_entidad, delete_entidad,
    contar_subentidades_activas,
    get_all_usuarios, create_usuario, update_usuario, delete_usuario,
    documento_existe, email_existe, username_existe,
    username_existe_otro, email_existe_otro,
    get_guias_caninos,
    get_all_caninos, get_canino_by_id, create_canino, update_canino, dar_baja_canino,
    get_all_vacunas_agrupadas, create_vacuna_canino, delete_vacuna_canino, update_historial_canino,
    get_sesiones_by_especialidad, count_sesiones_por_especialidad,
    create_sesion_curso, delete_sesion_curso,
    get_all_programas, get_programa_by_clave, create_programa, update_programa, delete_programa,
    get_sesion_by_id, get_vision_archivos, create_vision_archivo, delete_vision_archivo,
    get_all_vision_archivos,
    update_sesion_curso,
    get_user_by_username_o_email, crear_token_recuperacion,
    get_token_recuperacion_valido, usar_token_recuperacion,
)
from .correo import enviar_correo
from psycopg import errors as pg_errors
from .auth import (
    crear_sesion, es_sesion_valida, es_super_admin, cerrar_sesion,
    redirigir_si_no_autenticado, redirigir_si_no_super_admin,
)

# ---------------------------------------------------------------------------
# App y middleware
# ---------------------------------------------------------------------------

app = FastAPI()

app.add_middleware(
    SessionMiddleware,
    secret_key=os.getenv('SECRET_KEY', 'dev-secret-change-in-prod'),
    max_age=int(os.getenv('SESSION_LIFETIME', 3600)),
    same_site='lax',
    https_only=False,
)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
templates = Jinja2Templates(directory=os.path.join(BASE_DIR, 'templates'))

static_path = os.path.join(BASE_DIR, 'static')
if not os.path.exists(static_path):
    os.makedirs(static_path)

app.mount('/static', StaticFiles(directory=static_path), name='static')


# ---------------------------------------------------------------------------
# Helper: contexto base para todas las plantillas protegidas
# Pasa siempre la sesión para que layout.html pueda usarla en el sidebar
# ---------------------------------------------------------------------------

def ctx(request: Request, **extra):
    return {'request': request, 'session': request.session, **extra}


# ---------------------------------------------------------------------------
# Rutas públicas
# ---------------------------------------------------------------------------

@app.get('/', response_class=HTMLResponse)
async def root(request: Request):
    if es_sesion_valida(request):
        return RedirectResponse(url='/index', status_code=302)
    return RedirectResponse(url='/login', status_code=302)


@app.get('/login', response_class=HTMLResponse)
async def login_page(request: Request):
    if es_sesion_valida(request):
        return RedirectResponse(url='/index', status_code=302)
    return templates.TemplateResponse('login.html', {'request': request, 'error': None})


@app.post('/login', response_class=HTMLResponse)
async def login_post(request: Request,
                     username: str = Form(...),
                     password: str = Form(...)):
    user = verify_user(username, password)
    if not user:
        return templates.TemplateResponse(
            'login.html',
            {'request': request, 'error': 'Usuario o contraseña incorrectos'},
            status_code=401,
        )
    crear_sesion(request, user)
    return RedirectResponse(url='/index', status_code=303)


@app.get('/logout')
async def logout(request: Request):
    cerrar_sesion(request)
    return RedirectResponse(url='/login', status_code=302)


@app.get('/recuperar', response_class=HTMLResponse)
async def recover_page(request: Request):
    return templates.TemplateResponse('recuperar.html', {'request': request})


# ---------------------------------------------------------------------------
# Recuperación de contraseña por correo
#   1. POST /recuperar     -> genera un token y envía el enlace al correo
#   2. GET  /restablecer   -> formulario de nueva contraseña (si el token vale)
#   3. POST /restablecer   -> guarda la nueva contraseña y "quema" el token
# ---------------------------------------------------------------------------

MINUTOS_TOKEN = 30

# Mensaje igual exista o no el usuario: así nadie puede averiguar qué
# nombres de usuario o correos están registrados.
MSG_RECUPERAR_OK = ('Si los datos corresponden a un usuario activo, enviamos un enlace '
                    'a su correo registrado. Revise su bandeja (y la carpeta de spam).')


@app.post('/recuperar', response_class=HTMLResponse)
async def recover_post(request: Request, identificador: str = Form(...)):
    user = get_user_by_username_o_email(identificador.strip())

    if user and user.get('estado') == 'activo':
        # El token en claro solo viaja en el enlace; en la base va su SHA-256.
        token = secrets.token_urlsafe(32)
        token_hash = hashlib.sha256(token.encode('utf-8')).hexdigest()
        ip = request.client.host if request.client else None
        crear_token_recuperacion(user['id'], token_hash, MINUTOS_TOKEN, ip)

        app_url = os.getenv('APP_URL', str(request.base_url)).rstrip('/')
        enlace = f'{app_url}/restablecer?token={token}'
        datos = {
            'nombre': user['nombre'], 'nombre_usuario': user['nombre_usuario'],
            'enlace': enlace, 'minutos': MINUTOS_TOKEN,
        }
        # Versión con diseño (templates/emails/recuperacion.html)...
        html = templates.get_template('emails/recuperacion.html').render(**datos)
        # ...y versión en texto plano para clientes de correo que no muestran HTML.
        texto = (
            f"Hola {datos['nombre']},\n\n"
            f"Recibimos una solicitud para restablecer la contraseña de tu cuenta "
            f"({datos['nombre_usuario']}).\n\n"
            f"Crea una nueva contraseña aquí (vence en {MINUTOS_TOKEN} minutos):\n"
            f"{enlace}\n\n"
            f"Si no solicitaste este cambio, ignora este correo."
        )
        enviar_correo(user['email'], 'Restablece tu contraseña · CanGuard K9', texto, html)

    return templates.TemplateResponse('recuperar.html', {'request': request, 'ok': MSG_RECUPERAR_OK})


@app.get('/restablecer', response_class=HTMLResponse)
async def reset_page(request: Request, token: str = ''):
    token_hash = hashlib.sha256(token.encode('utf-8')).hexdigest()
    valido = bool(token) and get_token_recuperacion_valido(token_hash) is not None
    return templates.TemplateResponse('restablecer.html', {
        'request': request, 'token': token, 'valido': valido, 'error': None,
    })


@app.post('/restablecer', response_class=HTMLResponse)
async def reset_post(request: Request,
                     token: str = Form(...),
                     password: str = Form(...),
                     confirmar: str = Form(...)):
    token_hash = hashlib.sha256(token.encode('utf-8')).hexdigest()
    registro = get_token_recuperacion_valido(token_hash)

    def responder(error, valido=True):
        return templates.TemplateResponse('restablecer.html', {
            'request': request, 'token': token, 'valido': valido, 'error': error,
        }, status_code=400)

    if not registro:
        return responder(None, valido=False)
    if len(password) < 6:
        return responder('La contraseña debe tener al menos 6 caracteres.')
    if password != confirmar:
        return responder('Las contraseñas no coinciden.')

    usar_token_recuperacion(registro['id'], registro['id_usuario'], password)
    return templates.TemplateResponse('login.html', {
        'request': request, 'error': None,
        'ok': 'Contraseña actualizada. Ya puede iniciar sesión.',
    })


# ---------------------------------------------------------------------------
# Dashboard principal
# ---------------------------------------------------------------------------

@app.get('/index', response_class=HTMLResponse)
async def index_page(request: Request):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r

    es_sa  = es_super_admin(request)
    id_ent = None if es_sa else request.session.get('id_entidad')

    caninos          = get_all_caninos(id_entidad=id_ent)              or []
    guias            = get_guias_caninos(id_entidad=id_ent)            or []
    sesiones_por_esp = count_sesiones_por_especialidad(id_entidad=id_ent) or {}
    vision_archs     = get_all_vision_archivos(id_entidad=id_ent)      or []
    usuarios         = get_all_usuarios(id_entidad=id_ent)             or []

    total_sesiones  = sum(sesiones_por_esp.values()) if sesiones_por_esp else 0
    total_vision    = len(vision_archs)
    con_trans       = sum(1 for a in vision_archs if a.get('notas'))
    caninos_activos = sum(1 for c in caninos if (c.get('estado') or 'activo') != 'baja')
    caninos_baja    = len(caninos) - caninos_activos
    caninos_entren  = sum(1 for c in caninos if c.get('estado') == 'entrenamiento')

    pct_trans = round(con_trans * 100 / total_vision) if total_vision else 0

    stats = {
        'total_caninos':    len(caninos),
        'caninos_activos':  caninos_activos,
        'caninos_baja':     caninos_baja,
        'caninos_entren':   caninos_entren,
        'total_guias':      len(guias),
        'total_sesiones':   total_sesiones,
        'sesiones_por_esp': sesiones_por_esp,
        'total_vision':     total_vision,
        'con_transcripcion': con_trans,
        'pct_transcripcion': pct_trans,
        'total_usuarios':   len(usuarios),
    }
    if es_sa:
        entidades = get_all_entidades() or []
        stats['total_entidades'] = len(entidades)

    return templates.TemplateResponse('dashboard.html', ctx(request, stats=stats))


# ---------------------------------------------------------------------------
# Entidades  (solo super_admin)
# ---------------------------------------------------------------------------

@app.get('/entidades', response_class=HTMLResponse)
async def listar_entidades(request: Request):
    r = redirigir_si_no_super_admin(request)
    if r:
        return r
    entidades = get_all_entidades() or []
    return templates.TemplateResponse('entidades.html', ctx(
        request, entidades=entidades, modo='entidades', flash=request.session.pop('flash_entidad', None)))


@app.get('/subentidades', response_class=HTMLResponse)
async def listar_subentidades(request: Request, padre: Optional[int] = None):
    """Módulo de sub-entidades (CAI, estaciones, ...). Usa la misma plantilla que
    entidades en modo 'sub'. Si llega ?padre=ID, el formulario se abre con esa
    entidad principal ya seleccionada (botón de la lista de entidades)."""
    r = redirigir_si_no_super_admin(request)
    if r:
        return r
    entidades = get_all_entidades() or []
    return templates.TemplateResponse(
        'entidades.html', ctx(request, entidades=entidades, modo='sub', abrir_padre=padre,
                              flash=request.session.pop('flash_entidad', None)))


def _flash_entidad(request: Request, titulo: str, texto: str):
    """Guarda en la sesión el mensaje de éxito; la lista de entidades lo lee
    (y lo borra) al cargar, y lo muestra con SweetAlert."""
    request.session['flash_entidad'] = {'titulo': titulo, 'texto': texto}


def _guardar_logo(razon_social: str, archivo: Optional[UploadFile]) -> str | None:
    """Guarda el logo en static/images/{slug}/logo/ y retorna la ruta relativa."""
    if not archivo or not archivo.filename:
        return None
    slug = re.sub(r'[^a-z0-9]+', '_', razon_social.lower()).strip('_')
    ext  = os.path.splitext(archivo.filename)[1].lower()
    carpeta = os.path.join(BASE_DIR, 'static', 'images', slug, 'logo')
    os.makedirs(carpeta, exist_ok=True)
    ruta_disco = os.path.join(carpeta, f'logo{ext}')
    with open(ruta_disco, 'wb') as f:
        shutil.copyfileobj(archivo.file, f)
    return f'images/{slug}/logo/logo{ext}'   # relativo a /static/


@app.get('/RegistroEntidad', response_class=HTMLResponse)
async def ver_registro_entidad(request: Request):
    # Mantiene la ruta GET por compatibilidad; el flujo principal usa el modal en /entidades
    r = redirigir_si_no_super_admin(request)
    if r:
        return r
    return RedirectResponse(url='/entidades', status_code=302)


def _modo(volver: str) -> str:
    """Desde qué módulo se envió el formulario: 'sub' (Sub-entidades) o 'entidades'."""
    return 'sub' if volver == 'subentidades' else 'entidades'


def _url_modulo(volver: str) -> str:
    # Solo se aceptan estas dos rutas, para que el formulario no pueda redirigir a cualquier lado.
    return '/subentidades' if _modo(volver) == 'sub' else '/entidades'


def _error_entidades(request: Request, mensaje: str, volver: str = 'entidades'):
    """Vuelve a mostrar el módulo (entidades o sub-entidades) con un mensaje de error."""
    entidades = get_all_entidades() or []
    return templates.TemplateResponse(
        'entidades.html', ctx(request, entidades=entidades, error=mensaje, modo=_modo(volver)),
        status_code=400,
    )


def _id_padre(valor: str) -> int | None:
    """El <select> de entidad padre manda '' cuando es una entidad principal."""
    return int(valor) if valor and valor.isdigit() else None


def _error_jerarquia(id_padre: int | None, id_ent: int | None = None) -> str | None:
    """Solo hay dos niveles: Entidad principal -> Sub-entidad.
    Devuelve el mensaje de error si la combinación no es válida, o None si está bien."""
    if not id_padre:
        return None
    padre = get_entidad_by_id(id_padre)
    if not padre:
        return 'La entidad principal seleccionada no existe.'
    if padre.get('id_entidad_padre'):
        return 'Una sub-entidad no puede tener sub-entidades. Seleccione una entidad principal.'
    if id_ent and (id_padre == id_ent or contar_subentidades_activas(id_ent)):
        return 'Esta entidad tiene sub-entidades, por eso no puede convertirse en sub-entidad.'
    return None


def _avisar_entidad_creada(request: Request, id_ent: int):
    """Envía al correo administrativo de la entidad el aviso de que fue creada,
    con sus datos para que los verifique."""
    ent = get_entidad_by_id(id_ent)
    if not ent or not ent.get('email'):
        return
    padre = get_entidad_by_id(ent['id_entidad_padre']) if ent.get('id_entidad_padre') else None
    datos = {
        'razon_social': ent['razon_social'], 'nit_cif': ent['nit_cif'],
        'ciudad': ent['ciudad'], 'direccion': ent['direccion'], 'telefono': ent['telefono'],
        'nombre_padre': padre['razon_social'] if padre else None,
        # A quién escribir si los datos están mal: CONTACTO_ADMIN del .env, o el remitente.
        # parseaddr deja solo la dirección: 'CanGuard K9 <x@y.com>' -> 'x@y.com'.
        'contacto': parseaddr(os.getenv('CONTACTO_ADMIN') or os.getenv('SMTP_FROM') or '')[1],
        'app_url': os.getenv('APP_URL', str(request.base_url)).rstrip('/'),
    }
    html = templates.get_template('emails/entidad_creada.html').render(**datos)
    lineas = [
        f"Su entidad {datos['razon_social']} fue registrada en el sistema CanGuard K9.",
        "",
        "Verifique que los datos sean correctos:",
        f"  Razón social: {datos['razon_social']}",
        f"  NIT: {datos['nit_cif']}",
    ]
    if datos['nombre_padre']:
        lineas.append(f"  Depende de: {datos['nombre_padre']}")
    lineas += [
        f"  Ciudad: {datos['ciudad'] or '—'}",
        f"  Dirección: {datos['direccion'] or '—'}",
        f"  Teléfono: {datos['telefono'] or '—'}",
        "",
        "Si algún dato no es correcto, contacte al administrador de CanGuard"
        + (f": {datos['contacto']}" if datos['contacto'] else "."),
    ]
    enviar_correo(ent['email'], 'Su entidad fue registrada · CanGuard K9', '\n'.join(lineas), html)


@app.post('/RegistroEntidad', response_class=HTMLResponse)
async def crear_entidad_post(
    request:          Request,
    RazonSocial:      str        = Form(...),
    NitCif:           str        = Form(...),
    TelefonoEntidad:  str        = Form(''),
    PaisEntidad:      str        = Form('CO'),
    UbicacionEntidad: str        = Form(''),
    DireccionEntidad: str        = Form(''),
    EmailEntidad:     str        = Form(''),
    IdEntidadPadre:   str        = Form(''),
    Volver:           str        = Form('entidades'),
    Logo:             Optional[UploadFile] = File(None),
):
    r = redirigir_si_no_super_admin(request)
    if r:
        return r
    id_padre = _id_padre(IdEntidadPadre)
    # En el módulo Sub-entidades la entidad principal es obligatoria.
    if _modo(Volver) == 'sub' and not id_padre:
        return _error_entidades(request, 'Seleccione la entidad principal de la sub-entidad.', Volver)
    error = _error_jerarquia(id_padre)
    if error:
        return _error_entidades(request, error, Volver)
    logo_path = _guardar_logo(RazonSocial, Logo)
    try:
        id_ent = create_entidad(RazonSocial, NitCif, TelefonoEntidad,
                                PaisEntidad, UbicacionEntidad, DireccionEntidad, logo_path,
                                email=EmailEntidad.strip(), id_entidad_padre=id_padre)
    except pg_errors.UniqueViolation:
        # PostgreSQL lanza UniqueViolation cuando el NIT ya existe (restricción UqEntidadNit).
        return _error_entidades(request, 'Ya existe una entidad con ese NIT/CIF.', Volver)
    _avisar_entidad_creada(request, id_ent)
    tipo = 'Sub-entidad' if id_padre else 'Entidad'
    texto = f'{RazonSocial} quedó registrada en CanGuard.'
    if EmailEntidad.strip():
        texto += f' Se envió un aviso a {EmailEntidad.strip()}.'
    _flash_entidad(request, f'¡{tipo} creada con éxito!', texto)
    return RedirectResponse(url=_url_modulo(Volver), status_code=303)


@app.post('/entidades/{id_ent}/editar', response_class=HTMLResponse)
async def editar_entidad_post(
    request:          Request,
    id_ent:           int,
    RazonSocial:      str        = Form(...),
    NitCif:           str        = Form(...),
    TelefonoEntidad:  str        = Form(''),
    PaisEntidad:      str        = Form('CO'),
    UbicacionEntidad: str        = Form(''),
    DireccionEntidad: str        = Form(''),
    EmailEntidad:     str        = Form(''),
    IdEntidadPadre:   str        = Form(''),
    Volver:           str        = Form('entidades'),
    Logo:             Optional[UploadFile] = File(None),
):
    r = redirigir_si_no_super_admin(request)
    if r:
        return r
    id_padre = _id_padre(IdEntidadPadre)
    if _modo(Volver) == 'sub' and not id_padre:
        return _error_entidades(request, 'Seleccione la entidad principal de la sub-entidad.', Volver)
    error = _error_jerarquia(id_padre, id_ent)
    if error:
        return _error_entidades(request, error, Volver)
    logo_path = _guardar_logo(RazonSocial, Logo)
    try:
        update_entidad(id_ent, RazonSocial, NitCif, TelefonoEntidad,
                       PaisEntidad, UbicacionEntidad, DireccionEntidad, logo_path,
                       email=EmailEntidad.strip(), id_entidad_padre=id_padre)
    except pg_errors.UniqueViolation:
        return _error_entidades(request, 'Ya existe una entidad con ese NIT/CIF.', Volver)
    _flash_entidad(request, 'Cambios guardados', f'Los datos de {RazonSocial} se actualizaron correctamente.')
    return RedirectResponse(url=_url_modulo(Volver), status_code=303)


@app.post('/entidades/{id_ent}/eliminar')
async def eliminar_entidad_post(request: Request, id_ent: int, Volver: str = Form('entidades')):
    r = redirigir_si_no_super_admin(request)
    if r:
        return r
    # Si se desactiva una entidad con sub-entidades activas, estas quedarían
    # "huérfanas" y desaparecerían de la lista: primero hay que moverlas o eliminarlas.
    if contar_subentidades_activas(id_ent):
        return _error_entidades(request, 'No se puede eliminar: la entidad tiene sub-entidades activas. '
                                         'Elimínelas o cámbielas de entidad primero.', Volver)
    ent = get_entidad_by_id(id_ent)
    delete_entidad(id_ent)
    if ent:
        _flash_entidad(request, 'Entidad eliminada', f'{ent["razon_social"]} fue eliminada.')
    return RedirectResponse(url=_url_modulo(Volver), status_code=303)


# ---------------------------------------------------------------------------
# Usuarios
# ---------------------------------------------------------------------------

@app.get('/usuarios', response_class=HTMLResponse)
async def listar_usuarios(request: Request):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    if es_super_admin(request):
        usuarios = get_all_usuarios() or []
    else:
        usuarios = get_all_usuarios(id_entidad=request.session.get('id_entidad')) or []

    entidades = get_all_entidades() if es_super_admin(request) else []

    flash_nuevo = request.session.get('flash_nuevo_usuario')
    if flash_nuevo:
        del request.session['flash_nuevo_usuario']

    form_error = request.session.get('flash_form_error')
    if form_error:
        del request.session['flash_form_error']

    flash_edit_id = request.session.get('flash_edit_id')
    if flash_edit_id:
        del request.session['flash_edit_id']

    return templates.TemplateResponse('usuarios.html', ctx(
        request,
        usuarios=usuarios,
        entidades=entidades,
        flash_nuevo=flash_nuevo,
        form_error=form_error,
        flash_edit_id=flash_edit_id,
    ))


@app.get('/usuarios/nuevo', response_class=HTMLResponse)
async def nuevo_usuario_form(request: Request):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    return RedirectResponse(url='/usuarios', status_code=302)


@app.post('/usuarios/nuevo', response_class=HTMLResponse)
async def nuevo_usuario_post(
    request:         Request,
    nombre:          str        = Form(...),
    nombre_usuario:  str        = Form(...),
    documento:       str        = Form(...),
    email:           str        = Form(...),
    telefono:        str        = Form(''),
    password:        str        = Form(...),
    password2:       str        = Form(...),
    rol:             str        = Form('admin_entidad'),
    id_entidad:      str        = Form(''),
    tipo_documento:  str        = Form(''),
    nacionalidad:    str        = Form('CO'),
    doc_identidad:   Optional[UploadFile] = File(None),
    foto_perfil:     Optional[UploadFile] = File(None),
):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r

    def _err(msg):
        request.session['flash_form_error'] = msg
        return RedirectResponse(url='/usuarios', status_code=303)

    if password != password2:
        return _err('Las contraseñas no coinciden.')
    if len(password) < 6:
        return _err('La contraseña debe tener al menos 6 caracteres.')
    if not nombre_usuario or len(nombre_usuario) < 3:
        return _err('El nombre de usuario debe tener al menos 3 caracteres.')
    if username_existe(nombre_usuario):
        return _err(f'El nombre de usuario "{nombre_usuario}" ya está en uso.')
    if documento_existe(documento):
        return _err('Ese número de documento ya está registrado.')
    if email_existe(email):
        return _err('Ese correo ya está registrado.')
    if es_super_admin(request) and rol != 'super_admin' and not id_entidad:
        return _err('Debes seleccionar una entidad para este rol.')

    if not es_super_admin(request):
        id_entidad = str(request.session.get('id_entidad') or '')
        rol = 'admin_entidad'

    # Guardar foto de perfil
    foto_path = None
    if foto_perfil and foto_perfil.filename:
        ext = os.path.splitext(foto_perfil.filename)[1].lower()
        carpeta = os.path.join(BASE_DIR, 'static', 'images', 'usuarios', nombre_usuario)
        os.makedirs(carpeta, exist_ok=True)
        with open(os.path.join(carpeta, f'foto{ext}'), 'wb') as fd:
            shutil.copyfileobj(foto_perfil.file, fd)
        foto_path = f'images/usuarios/{nombre_usuario}/foto{ext}'

    # Guardar documento de identidad
    doc_path = None
    if doc_identidad and doc_identidad.filename:
        ext = os.path.splitext(doc_identidad.filename)[1].lower()
        carpeta = os.path.join(BASE_DIR, 'static', 'docs', documento)
        os.makedirs(carpeta, exist_ok=True)
        with open(os.path.join(carpeta, f'doc{ext}'), 'wb') as fd:
            shutil.copyfileobj(doc_identidad.file, fd)
        doc_path = f'docs/{documento}/doc{ext}'

    id_ent_int = int(id_entidad) if id_entidad else None
    create_usuario(nombre, nombre_usuario, documento, email, telefono, password, rol,
                   id_ent_int, tipo_documento or None, doc_path, nacionalidad, foto_path)

    # Flash: super_admin ve las credenciales en texto claro una sola vez
    if es_super_admin(request):
        request.session['flash_nuevo_usuario'] = {
            'nombre':         nombre,
            'nombre_usuario': nombre_usuario,
            'password':       password,
        }

    return RedirectResponse(url='/usuarios', status_code=303)


@app.post('/usuarios/{id_usr}/editar', response_class=HTMLResponse)
async def editar_usuario_post(
    request:         Request,
    id_usr:          int,
    nombre:          str        = Form(...),
    nombre_usuario:  str        = Form(...),
    email:           str        = Form(...),
    telefono:        str        = Form(''),
    password:        str        = Form(''),
    password2:       str        = Form(''),
    rol:             str        = Form(''),
    id_entidad:      str        = Form(''),
    foto_perfil:     Optional[UploadFile] = File(None),
):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r

    def _err(msg):
        request.session['flash_form_error'] = msg
        request.session['flash_edit_id']    = id_usr
        return RedirectResponse(url='/usuarios', status_code=303)

    if password and password != password2:
        return _err('Las contraseñas no coinciden.')
    if password and len(password) < 6:
        return _err('La contraseña debe tener al menos 6 caracteres.')
    if username_existe_otro(nombre_usuario, id_usr):
        return _err(f'El nombre de usuario "{nombre_usuario}" ya está en uso.')
    if email_existe_otro(email, id_usr):
        return _err('Ese correo ya está registrado por otro usuario.')

    foto_path = None
    if foto_perfil and foto_perfil.filename:
        ext = os.path.splitext(foto_perfil.filename)[1].lower()
        carpeta = os.path.join(BASE_DIR, 'static', 'images', 'usuarios', nombre_usuario)
        os.makedirs(carpeta, exist_ok=True)
        with open(os.path.join(carpeta, f'foto{ext}'), 'wb') as fd:
            shutil.copyfileobj(foto_perfil.file, fd)
        foto_path = f'images/usuarios/{nombre_usuario}/foto{ext}'

    rol_final     = rol if (rol and es_super_admin(request)) else None
    id_ent_final  = (int(id_entidad) if id_entidad else None) if es_super_admin(request) else False

    update_usuario(
        id_usr, nombre, nombre_usuario, email, telefono,
        password        = password   if password   else None,
        rol             = rol_final,
        id_entidad      = id_ent_final,
        foto_perfil_path= foto_path,
    )
    return RedirectResponse(url='/usuarios', status_code=303)


@app.post('/usuarios/{id_usr}/eliminar')
async def eliminar_usuario_post(request: Request, id_usr: int):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    delete_usuario(id_usr)
    return RedirectResponse(url='/usuarios', status_code=303)


# ---------------------------------------------------------------------------
# Caninos
# ---------------------------------------------------------------------------

def _guardar_foto_canino(nombre: str, archivo: Optional[UploadFile]) -> str | None:
    if not archivo or not archivo.filename:
        return None
    slug = re.sub(r'[^a-z0-9]+', '_', nombre.lower()).strip('_')
    ext  = os.path.splitext(archivo.filename)[1].lower()
    carpeta = os.path.join(BASE_DIR, 'static', 'images', 'caninos', slug)
    os.makedirs(carpeta, exist_ok=True)
    ruta = os.path.join(carpeta, f'foto{ext}')
    with open(ruta, 'wb') as f:
        shutil.copyfileobj(archivo.file, f)
    return f'images/caninos/{slug}/foto{ext}'


@app.get('/caninos', response_class=HTMLResponse)
async def listar_caninos(request: Request):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    if es_super_admin(request):
        caninos   = get_all_caninos(incluir_baja=True) or []
        entidades = get_all_entidades() or []
        guias     = get_guias_caninos() or []
    else:
        id_ent    = request.session.get('id_entidad')
        caninos   = get_all_caninos(id_entidad=id_ent, incluir_baja=True) or []
        entidades = []
        guias     = get_guias_caninos(id_entidad=id_ent) or []

    flash_ok = request.session.get('flash_canino_ok')
    if flash_ok:
        del request.session['flash_canino_ok']

    form_error = request.session.get('flash_canino_err')
    if form_error:
        del request.session['flash_canino_err']

    flash_edit_id = request.session.get('flash_canino_edit_id')
    if flash_edit_id:
        del request.session['flash_canino_edit_id']

    return templates.TemplateResponse('caninos.html', ctx(
        request,
        caninos=caninos,
        entidades=entidades,
        guias=guias,
        flash_ok=flash_ok,
        form_error=form_error,
        flash_edit_id=flash_edit_id,
    ))


@app.post('/caninos/nuevo', response_class=HTMLResponse)
async def nuevo_canino_post(
    request:            Request,
    nombre:             str  = Form(...),
    raza:               str  = Form(...),
    genero:             str  = Form(...),
    especialidad:       str  = Form(...),
    estado:             str  = Form('activo'),
    id_entidad:         str  = Form(''),
    id_guia:            str  = Form(''),
    fecha_nacimiento:   str  = Form(''),
    chip_numero:        str  = Form(''),
    num_registro:       str  = Form(''),
    color_pelaje:       str  = Form(''),
    peso_kg:            str  = Form(''),
    talla_cm:           str  = Form(''),
    fecha_ingreso:      str  = Form(''),
    procedencia:        str  = Form(''),
    nombre_vacuna:      str  = Form(''),
    ultima_vacunacion:  str  = Form(''),
    proxima_vacunacion: str  = Form(''),
    observaciones:      str  = Form(''),
    foto_canino:        Optional[UploadFile] = File(None),
):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r

    def _err(msg):
        request.session['flash_canino_err'] = msg
        return RedirectResponse(url='/caninos', status_code=303)

    if not es_super_admin(request):
        id_entidad = str(request.session.get('id_entidad') or '')

    if not id_entidad:
        return _err('Debes seleccionar una entidad.')

    foto_path = _guardar_foto_canino(nombre, foto_canino)

    create_canino(
        nombre=nombre, raza=raza, genero=genero,
        especialidad=especialidad, estado=estado,
        id_entidad=int(id_entidad),
        id_guia=int(id_guia) if id_guia else None,
        fecha_nacimiento=fecha_nacimiento     or None,
        chip_numero=chip_numero               or None,
        num_registro=num_registro             or None,
        color_pelaje=color_pelaje             or None,
        peso_kg=float(peso_kg)   if peso_kg   else None,
        talla_cm=int(talla_cm)   if talla_cm  else None,
        foto_path=foto_path,
        fecha_ingreso=fecha_ingreso           or None,
        procedencia=procedencia               or None,
        nombre_vacuna=nombre_vacuna           or None,
        ultima_vacunacion=ultima_vacunacion   or None,
        proxima_vacunacion=proxima_vacunacion or None,
        observaciones=observaciones           or None,
    )
    request.session['flash_canino_ok'] = f'Canino "{nombre}" registrado correctamente.'
    return RedirectResponse(url='/caninos', status_code=303)


@app.post('/caninos/{id_can}/editar', response_class=HTMLResponse)
async def editar_canino_post(
    request:            Request,
    id_can:             int,
    nombre:             str  = Form(...),
    raza:               str  = Form(...),
    genero:             str  = Form(...),
    especialidad:       str  = Form(...),
    estado:             str  = Form('activo'),
    id_entidad:         str  = Form(''),
    id_guia:            str  = Form(''),
    fecha_nacimiento:   str  = Form(''),
    chip_numero:        str  = Form(''),
    num_registro:       str  = Form(''),
    color_pelaje:       str  = Form(''),
    peso_kg:            str  = Form(''),
    talla_cm:           str  = Form(''),
    fecha_ingreso:      str  = Form(''),
    procedencia:        str  = Form(''),
    nombre_vacuna:      str  = Form(''),
    ultima_vacunacion:  str  = Form(''),
    proxima_vacunacion: str  = Form(''),
    observaciones:      str  = Form(''),
    foto_canino:        Optional[UploadFile] = File(None),
):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r

    def _err(msg):
        request.session['flash_canino_err']     = msg
        request.session['flash_canino_edit_id'] = id_can
        return RedirectResponse(url='/caninos', status_code=303)

    if not es_super_admin(request):
        id_entidad = str(request.session.get('id_entidad') or '')

    foto_path = _guardar_foto_canino(nombre, foto_canino)

    update_canino(
        id_canino=id_can,
        nombre=nombre, raza=raza, genero=genero,
        especialidad=especialidad, estado=estado,
        id_entidad=int(id_entidad) if id_entidad else None,
        id_guia=int(id_guia) if id_guia else None,
        fecha_nacimiento=fecha_nacimiento     or None,
        chip_numero=chip_numero               or None,
        num_registro=num_registro             or None,
        color_pelaje=color_pelaje             or None,
        peso_kg=float(peso_kg)   if peso_kg   else None,
        talla_cm=int(talla_cm)   if talla_cm  else None,
        foto_path=foto_path,
        fecha_ingreso=fecha_ingreso           or None,
        procedencia=procedencia               or None,
        nombre_vacuna=nombre_vacuna           or None,
        ultima_vacunacion=ultima_vacunacion   or None,
        proxima_vacunacion=proxima_vacunacion or None,
        observaciones=observaciones           or None,
    )
    return RedirectResponse(url='/caninos', status_code=303)


@app.post('/caninos/{id_can}/eliminar')
async def dar_baja_canino_post(
    request:           Request,
    id_can:            int,
    motivo_baja:       str = Form(...),
    justificacion_baja: str = Form(...),
    fecha_baja:        str = Form(...),
    carta_baja:        Optional[UploadFile] = File(None),
):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r

    carta_path = None
    if carta_baja and carta_baja.filename:
        import time as _time
        ext     = os.path.splitext(carta_baja.filename)[1].lower()
        carpeta = os.path.join(BASE_DIR, 'static', 'docs', 'bajas', str(id_can))
        os.makedirs(carpeta, exist_ok=True)
        fname   = f'carta_baja_{int(_time.time())}{ext}'
        with open(os.path.join(carpeta, fname), 'wb') as f:
            shutil.copyfileobj(carta_baja.file, f)
        carta_path = f'docs/bajas/{id_can}/{fname}'

    dar_baja_canino(id_can, motivo_baja, justificacion_baja, fecha_baja, carta_path)
    return RedirectResponse(url='/caninos', status_code=303)


# ---------------------------------------------------------------------------
# Salud Canina
# ---------------------------------------------------------------------------

def _estado_salud(vacunas: list) -> str:
    """Calcula el estado de salud según las fechas de vencimiento."""
    if not vacunas:
        return 'sin_registro'
    hoy    = date.today()
    pronto = hoy + timedelta(days=30)
    estado = 'al_dia'
    for v in vacunas:
        fv = v.get('fecha_vencimiento')
        if fv:
            if fv < hoy:
                return 'vencida'
            if fv <= pronto and estado != 'vencida':
                estado = 'por_vencer'
    return estado


@app.get('/salud-canina', response_class=HTMLResponse)
async def salud_canina_page(request: Request):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    if es_super_admin(request):
        caninos   = get_all_caninos() or []
        entidades = get_all_entidades() or []
    else:
        caninos   = get_all_caninos(id_entidad=request.session.get('id_entidad')) or []
        entidades = []

    ids       = [c['id'] for c in caninos]
    vac_map   = get_all_vacunas_agrupadas(ids)
    salud_map = {c['id']: _estado_salud(vac_map.get(c['id'], [])) for c in caninos}

    hoy = date.today()
    alertas = []
    nombre_map = {c['id']: c['nombre'] for c in caninos}
    for id_can, vacunas in vac_map.items():
        for v in vacunas:
            fv = v.get('fecha_vencimiento')
            if fv:
                dias = (fv - hoy).days
                if 0 <= dias <= 3:
                    alertas.append({
                        'canino': nombre_map.get(id_can, ''),
                        'vacuna': v.get('nombre_vacuna', ''),
                        'dias':   dias,
                        'fecha':  str(fv),
                    })

    flash_ok = request.session.get('flash_salud_ok')
    if flash_ok:
        del request.session['flash_salud_ok']

    return templates.TemplateResponse('salud.html', ctx(
        request,
        caninos=caninos,
        entidades=entidades,
        vac_map=vac_map,
        salud_map=salud_map,
        flash_ok=flash_ok,
        today=hoy,
        alertas=alertas,
    ))


@app.post('/salud-canina/{id_can}/vacuna', response_class=HTMLResponse)
async def agregar_vacuna_post(
    request:           Request,
    id_can:            int,
    nombre_vacuna:     str = Form(...),
    fecha_aplicacion:  str = Form(...),
    fecha_vencimiento: str = Form(''),
    veterinario:       str = Form(''),
    notas:             str = Form(''),
    comprobante:       Optional[UploadFile] = File(None),
):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r

    comprobante_path = None
    if comprobante and comprobante.filename:
        import time as _time
        ext     = os.path.splitext(comprobante.filename)[1].lower()
        carpeta = os.path.join(BASE_DIR, 'static', 'docs', 'vacunas', str(id_can))
        os.makedirs(carpeta, exist_ok=True)
        fname   = f'vac_{int(_time.time())}{ext}'
        with open(os.path.join(carpeta, fname), 'wb') as f:
            shutil.copyfileobj(comprobante.file, f)
        comprobante_path = f'docs/vacunas/{id_can}/{fname}'

    create_vacuna_canino(
        id_can, nombre_vacuna, fecha_aplicacion,
        fecha_vencimiento or None, veterinario or None, notas or None,
        comprobante_path,
    )
    request.session['flash_salud_ok'] = 'Vacuna registrada correctamente.'
    return RedirectResponse(url='/salud-canina', status_code=303)


@app.post('/salud-canina/vacunas/{id_vac}/eliminar')
async def eliminar_vacuna_post(request: Request, id_vac: int):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    delete_vacuna_canino(id_vac)
    return RedirectResponse(url='/salud-canina', status_code=303)


@app.post('/salud-canina/{id_can}/historial', response_class=HTMLResponse)
async def subir_historial_post(
    request:   Request,
    id_can:    int,
    historial: UploadFile = File(...),
):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    if historial and historial.filename:
        ext     = os.path.splitext(historial.filename)[1].lower()
        carpeta = os.path.join(BASE_DIR, 'static', 'docs', 'caninos', str(id_can))
        os.makedirs(carpeta, exist_ok=True)
        ruta    = os.path.join(carpeta, f'historial{ext}')
        with open(ruta, 'wb') as f:
            shutil.copyfileobj(historial.file, f)
        update_historial_canino(id_can, f'docs/caninos/{id_can}/historial{ext}')
        request.session['flash_salud_ok'] = 'Historial clínico subido correctamente.'
    return RedirectResponse(url='/salud-canina', status_code=303)


# ---------------------------------------------------------------------------
# Notificaciones
# ---------------------------------------------------------------------------

def _calcular_alertas(caninos: list, vac_map: dict) -> list:
    """Devuelve vacunas que vencen en ≤ 3 días o ya vencidas."""
    hoy    = date.today()
    nombre = {c['id']: c['nombre'] for c in caninos}
    items  = []
    for id_can, vacunas in vac_map.items():
        for v in vacunas:
            fv = v.get('fecha_vencimiento')
            if fv:
                dias = (fv - hoy).days
                if dias <= 3:
                    items.append({
                        'canino': nombre.get(id_can, ''),
                        'vacuna': v.get('nombre_vacuna', ''),
                        'dias':   dias,
                        'fecha':  str(fv),
                        'vencida': dias < 0,
                    })
    items.sort(key=lambda x: x['dias'])
    return items


@app.get('/api/notificaciones/count')
async def api_notif_count(request: Request):
    if not request.session.get('usuario_id'):
        return JSONResponse({'count': 0, 'items': []})
    if es_super_admin(request):
        caninos = get_all_caninos() or []
    else:
        caninos = get_all_caninos(id_entidad=request.session.get('id_entidad')) or []
    ids     = [c['id'] for c in caninos]
    vac_map = get_all_vacunas_agrupadas(ids)
    items   = _calcular_alertas(caninos, vac_map)
    return JSONResponse({'count': len(items), 'items': items})


@app.get('/notificaciones', response_class=HTMLResponse)
async def notificaciones_page(request: Request):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    if es_super_admin(request):
        caninos   = get_all_caninos() or []
        entidades = get_all_entidades() or []
    else:
        caninos   = get_all_caninos(id_entidad=request.session.get('id_entidad')) or []
        entidades = []
    ids     = [c['id'] for c in caninos]
    vac_map = get_all_vacunas_agrupadas(ids)
    alertas = _calcular_alertas(caninos, vac_map)
    return templates.TemplateResponse('notificaciones.html', ctx(
        request,
        alertas=alertas,
        today=date.today(),
    ))


# ---------------------------------------------------------------------------
# Otras rutas protegidas
# ---------------------------------------------------------------------------

@app.get('/gestion-guias', response_class=HTMLResponse)
async def ver_registro_guias(request: Request):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    if es_super_admin(request):
        guias     = get_guias_caninos() or []
        entidades = get_all_entidades() or []
    else:
        guias     = get_guias_caninos(id_entidad=request.session.get('id_entidad')) or []
        entidades = []
    return templates.TemplateResponse('registro.html', ctx(request, guias=guias, entidades=entidades))


@app.get('/geolocalizacion', response_class=HTMLResponse)
async def geolocalizacion(request: Request):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    return templates.TemplateResponse('geolocalizacion.html', ctx(request))


@app.get('/vision-artificial', response_class=HTMLResponse)
async def vision_artificial(request: Request):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    id_ent = None if es_super_admin(request) else request.session.get('id_entidad')
    archivos  = get_all_vision_archivos(id_ent)
    flash_ok  = request.session.get('flash_vision_ok')
    if flash_ok:
        del request.session['flash_vision_ok']
    return templates.TemplateResponse('vision_artificial.html', ctx(
        request, archivos=archivos, flash_ok=flash_ok,
    ))


ESPECIALIDADES = [
    {'key': 'drogas',      'nombre': 'Detección de Drogas',       'icono': 'bi-eyedropper-fill',      'color': '#ef4444'},
    {'key': 'explosivos',  'nombre': 'Detección de Explosivos',   'icono': 'bi-lightning-charge-fill', 'color': '#f97316'},
    {'key': 'rescate',     'nombre': 'Búsqueda y Rescate',        'icono': 'bi-life-preserver',        'color': '#22c55e'},
    {'key': 'seguridad',   'nombre': 'Seguridad y Protección',    'icono': 'bi-shield-fill',           'color': '#3b82f6'},
    {'key': 'rastreo',     'nombre': 'Rastreo y Seguimiento',     'icono': 'bi-compass-fill',          'color': '#8b5cf6'},
    {'key': 'personas',    'nombre': 'Detección de Personas',     'icono': 'bi-people-fill',           'color': '#06b6d4'},
    {'key': 'multiple',    'nombre': 'Múltiple / Integral',       'icono': 'bi-collection-fill',       'color': '#b8860b'},
]


@app.get('/cursos', response_class=HTMLResponse)
async def cursos_page(request: Request):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    id_ent = None if es_super_admin(request) else request.session.get('id_entidad')

    programas = get_all_programas()
    conteos   = count_sesiones_por_especialidad(id_ent)

    flash_ok = request.session.get('flash_curso_ok')
    if flash_ok:
        del request.session['flash_curso_ok']

    return templates.TemplateResponse('cursos.html', ctx(
        request,
        programas=programas,
        conteos=conteos,
        flash_ok=flash_ok,
    ))


@app.get('/cursos/{esp}/sesiones', response_class=HTMLResponse)
async def cursos_sesiones_page(request: Request, esp: str):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    id_ent = None if es_super_admin(request) else request.session.get('id_entidad')
    curso  = get_programa_by_clave(esp)
    if not curso:
        return RedirectResponse('/cursos', status_code=303)

    sesiones  = get_sesiones_by_especialidad(esp, id_ent) or []
    if es_super_admin(request):
        caninos   = get_all_caninos() or []
        entidades = get_all_entidades() or []
        guias     = get_guias_caninos() or []
    else:
        caninos   = get_all_caninos(id_entidad=id_ent) or []
        entidades = []
        guias     = get_guias_caninos(id_entidad=id_ent) or []

    flash_ok = request.session.get('flash_curso_ok')
    if flash_ok:
        del request.session['flash_curso_ok']

    return templates.TemplateResponse('cursos_sesiones.html', ctx(
        request,
        curso=curso,
        sesiones=sesiones,
        caninos=caninos,
        entidades=entidades,
        guias=guias,
        flash_ok=flash_ok,
    ))


@app.post('/cursos/programa/nuevo')
async def crear_programa_post(
    request:       Request,
    nombre:        str = Form(...),
    descripcion:   str = Form(''),
    icono:         str = Form('bi-mortarboard-fill'),
    color:         str = Form('#b8860b'),
    imagen_fondo:  Optional[UploadFile] = File(None),
):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    if not es_super_admin(request):
        return RedirectResponse('/cursos', status_code=303)
    import time as _time
    clave = re.sub(r'[^a-z0-9_]', '', re.sub(r'\s+', '_', nombre.lower().strip()))[:50] or f'prog_{int(_time.time())}'
    img_path = None
    if imagen_fondo and imagen_fondo.filename:
        ext     = os.path.splitext(imagen_fondo.filename)[1].lower()
        carpeta = os.path.join(BASE_DIR, 'static', 'programas')
        os.makedirs(carpeta, exist_ok=True)
        fname   = f'prog_{int(_time.time())}{ext}'
        with open(os.path.join(carpeta, fname), 'wb') as f:
            shutil.copyfileobj(imagen_fondo.file, f)
        img_path = f'programas/{fname}'
    create_programa(clave, nombre, descripcion or None, icono, color, img_path)
    request.session['flash_curso_ok'] = f'Programa "{nombre}" creado correctamente.'
    return RedirectResponse('/cursos', status_code=303)


@app.post('/cursos/programa/{id_prog}/editar')
async def editar_programa_post(
    request:      Request,
    id_prog:      int,
    nombre:       str = Form(...),
    descripcion:  str = Form(''),
    icono:        str = Form('bi-mortarboard-fill'),
    color:        str = Form('#b8860b'),
    imagen_fondo: Optional[UploadFile] = File(None),
):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    if not es_super_admin(request):
        return RedirectResponse('/cursos', status_code=303)
    img_path = None
    if imagen_fondo and imagen_fondo.filename:
        import time as _time
        ext     = os.path.splitext(imagen_fondo.filename)[1].lower()
        carpeta = os.path.join(BASE_DIR, 'static', 'programas')
        os.makedirs(carpeta, exist_ok=True)
        fname   = f'prog_{int(_time.time())}{ext}'
        with open(os.path.join(carpeta, fname), 'wb') as f:
            shutil.copyfileobj(imagen_fondo.file, f)
        img_path = f'programas/{fname}'
    update_programa(id_prog, nombre, descripcion or None, icono, color, img_path)
    request.session['flash_curso_ok'] = f'Programa "{nombre}" actualizado correctamente.'
    return RedirectResponse('/cursos', status_code=303)


@app.post('/cursos/programa/{id_prog}/eliminar')
async def eliminar_programa_post(request: Request, id_prog: int):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    if not es_super_admin(request):
        return RedirectResponse('/cursos', status_code=303)
    delete_programa(id_prog)
    request.session['flash_curso_ok'] = 'Programa eliminado correctamente.'
    return RedirectResponse('/cursos', status_code=303)


@app.post('/cursos/{esp}/sesion')
async def crear_sesion_post(
    request:      Request,
    esp:          str,
    id_canino:    int = Form(...),
    id_entidad:   str = Form(''),
    fecha:        str = Form(...),
    resultado:    str = Form(''),
    instructor:   str = Form(''),
    observaciones: str = Form(''),
):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    id_ent_final = int(id_entidad) if id_entidad else request.session.get('id_entidad')
    create_sesion_curso(esp, id_canino, id_ent_final, fecha,
                        resultado or None, observaciones or None, instructor or None)
    request.session['flash_curso_ok'] = 'Sesión registrada correctamente.'
    return RedirectResponse(f'/cursos/{esp}/sesiones', status_code=303)


@app.post('/cursos/sesiones/{id_ses}/editar')
async def editar_sesion_post(
    request:      Request,
    id_ses:       int,
    fecha:        str = Form(...),
    resultado:    str = Form(''),
    instructor:   str = Form(''),
    observaciones: str = Form(''),
):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    update_sesion_curso(id_ses, fecha, resultado, observaciones, instructor)
    request.session['flash_curso_ok'] = 'Sesión actualizada correctamente.'
    referer = request.headers.get('referer', '/cursos')
    return RedirectResponse(referer, status_code=303)


@app.post('/cursos/sesiones/{id_ses}/eliminar')
async def eliminar_sesion_post(request: Request, id_ses: int):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    if request.session.get('rol') != 'super_admin':
        return RedirectResponse(request.headers.get('referer', '/cursos'), status_code=303)
    delete_sesion_curso(id_ses)
    referer = request.headers.get('referer', '/cursos')
    return RedirectResponse(referer, status_code=303)


# ---------------------------------------------------------------------------
# Visión Artificial por sesión
# ---------------------------------------------------------------------------

@app.get('/cursos/{esp}/sesiones/{id_ses}/vision', response_class=HTMLResponse)
async def vision_page(request: Request, esp: str, id_ses: int):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    sesion   = get_sesion_by_id(id_ses)
    if not sesion:
        return RedirectResponse(f'/cursos/{esp}/sesiones', status_code=303)
    curso    = get_programa_by_clave(esp)
    archivos = get_vision_archivos(id_ses)
    flash_ok = request.session.get('flash_vision_ok')
    if flash_ok:
        del request.session['flash_vision_ok']
    openai_ok = bool(os.getenv('OPENAI_API_KEY', ''))
    fresh  = bool(request.query_params.get('fresh'))
    review = bool(request.query_params.get('review'))
    return templates.TemplateResponse('cursos_vision.html', ctx(
        request, sesion=sesion, curso=curso, archivos=archivos,
        flash_ok=flash_ok, openai_ok=openai_ok, fresh=fresh, review=review,
    ))


@app.post('/cursos/{esp}/sesiones/{id_ses}/vision')
async def vision_upload(
    request:  Request,
    esp:      str,
    id_ses:   int,
    archivo:  UploadFile = File(...),
    notas:    str = Form(''),
):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    import time as _time
    ext     = os.path.splitext(archivo.filename)[1].lower()
    tipo    = 'video' if ext in ('.mp4', '.mov', '.avi', '.webm', '.ogg', '.m4a', '.mp3', '.wav') else 'imagen'
    carpeta = os.path.join(BASE_DIR, 'static', 'vision', str(id_ses))
    os.makedirs(carpeta, exist_ok=True)
    fname     = f'{int(_time.time())}{ext}'
    ruta_disk = os.path.join(carpeta, fname)
    with open(ruta_disk, 'wb') as f:
        shutil.copyfileobj(archivo.file, f)

    transcripcion = notas or ''
    if tipo in ('video', 'audio'):
        api_key = os.getenv('OPENAI_API_KEY', '')
        if api_key:
            try:
                from openai import OpenAI as _OAI
                with open(ruta_disk, 'rb') as af:
                    tr = _OAI(api_key=api_key).audio.transcriptions.create(
                        model='whisper-1', file=af, language='es')
                transcripcion = tr.text
            except Exception:
                try:
                    transcripcion = _transcribir_local(ruta_disk)
                except Exception:
                    pass
        else:
            try:
                transcripcion = _transcribir_local(ruta_disk)
            except Exception:
                pass

    create_vision_archivo(id_ses, f'vision/{id_ses}/{fname}', tipo, transcripcion or None)
    request.session['flash_vision_ok'] = 'Archivo cargado y transcrito correctamente.' if transcripcion else 'Archivo cargado.'
    return RedirectResponse(f'/cursos/{esp}/sesiones/{id_ses}/vision?review=1', status_code=303)


def _transcribir_local(ruta: str) -> str:
    """Transcribe usando Whisper local.
    Extrae el audio con imageio_ffmpeg directamente (sin depender del PATH)
    y pasa el array numpy a Whisper, evitando su load_audio interno."""
    import tempfile, subprocess, wave
    import numpy as np
    import whisper as _whisper

    # Obtener ffmpeg embebido en el venv (no el del PATH del sistema)
    try:
        import imageio_ffmpeg
        ffmpeg_exe = imageio_ffmpeg.get_ffmpeg_exe()
    except Exception:
        ffmpeg_exe = 'ffmpeg'

    tmp_fd, tmp_wav = tempfile.mkstemp(suffix='.wav')
    os.close(tmp_fd)
    try:
        # Extraer audio a WAV mono 16kHz
        cmd = [
            ffmpeg_exe, '-y', '-i', ruta,
            '-ar', '16000', '-ac', '1', '-f', 'wav',
            tmp_wav, '-loglevel', 'error'
        ]
        proc = subprocess.run(cmd, capture_output=True, timeout=300)
        if proc.returncode != 0:
            raise RuntimeError(proc.stderr.decode('utf-8', errors='replace') or 'Error extrayendo audio con ffmpeg')

        # Leer WAV como numpy float32 (Whisper acepta arrays directamente)
        with wave.open(tmp_wav, 'rb') as wf:
            raw = wf.readframes(wf.getnframes())
        audio = np.frombuffer(raw, dtype=np.int16).astype(np.float32) / 32768.0

        model = _whisper.load_model('base')
        result = model.transcribe(audio, language='es', fp16=False)
        return result['text']
    finally:
        if os.path.exists(tmp_wav):
            try:
                os.remove(tmp_wav)
            except Exception:
                pass


@app.post('/vision/{id_arch}/transcribir-ajax')
async def transcribir_ajax(request: Request, id_arch: int):
    r = redirigir_si_no_autenticado(request)
    if r:
        return JSONResponse({'ok': False, 'error': 'No autenticado'}, status_code=401)
    from .models import _query as _mq
    arch = _mq(
        "SELECT ArchivoPath AS archivo_path FROM RegistroDocumental WHERE IdRegistroDocumental=%s",
        (id_arch,), one=True
    )
    if not arch:
        return JSONResponse({'ok': False, 'error': 'Archivo no encontrado'})
    ruta = os.path.join(BASE_DIR, 'static', arch['archivo_path'].replace('/', os.sep))
    texto = None
    api_error = None

    # Intento 1: OpenAI API (si hay clave configurada)
    api_key = os.getenv('OPENAI_API_KEY', '')
    if api_key:
        try:
            from openai import OpenAI as _OAI
            with open(ruta, 'rb') as f:
                tr = _OAI(api_key=api_key).audio.transcriptions.create(
                    model='whisper-1', file=f, language='es')
            texto = tr.text
        except Exception as e:
            api_error = str(e)

    # Intento 2: Whisper local (fallback gratuito)
    if texto is None:
        try:
            texto = _transcribir_local(ruta)
        except Exception as e2:
            msg = api_error or ''
            if msg:
                msg += ' | '
            msg += 'Whisper local: ' + str(e2)
            return JSONResponse({'ok': False, 'error': msg})

    _mq("UPDATE RegistroDocumental SET Transcripcion=%s WHERE IdRegistroDocumental=%s", (texto, id_arch), write=True)
    return JSONResponse({'ok': True, 'texto': texto})


@app.post('/vision/{id_arch}/notas')
async def vision_guardar_notas(request: Request, id_arch: int, texto: str = Form('')):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    from .models import _query as _mq
    _mq("UPDATE RegistroDocumental SET Transcripcion=%s WHERE IdRegistroDocumental=%s", (texto or None, id_arch), write=True)
    request.session['flash_vision_ok'] = 'Transcripción guardada.'
    referer = request.headers.get('referer', '/cursos')
    sep = '&' if '?' in referer else '?'
    return RedirectResponse(referer + sep + 'fresh=1', status_code=303)


@app.post('/vision/{id_arch}/eliminar')
async def vision_eliminar(request: Request, id_arch: int):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    from .models import _query as _mq
    rows = _mq("SELECT ArchivoPath AS archivo_path FROM RegistroDocumental WHERE IdRegistroDocumental=%s", (id_arch,))
    if rows:
        ruta_disco = os.path.join(BASE_DIR, 'static', rows[0]['archivo_path'])
        try:
            if os.path.exists(ruta_disco):
                os.remove(ruta_disco)
            carpeta = os.path.dirname(ruta_disco)
            if os.path.isdir(carpeta) and not os.listdir(carpeta):
                os.rmdir(carpeta)
        except Exception:
            pass
    delete_vision_archivo(id_arch)
    referer = request.headers.get('referer', '/cursos')
    return RedirectResponse(referer, status_code=303)


@app.post('/cursos/{esp}/sesiones/{id_ses}/transcribir')
async def transcribir_vision(request: Request, esp: str, id_ses: int,
                              id_arch: int = Form(...)):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    api_key = os.getenv('OPENAI_API_KEY', '')
    if not api_key:
        request.session['flash_vision_ok'] = 'ERROR: Configura OPENAI_API_KEY en el archivo .env para usar la transcripción.'
        return RedirectResponse(f'/cursos/{esp}/sesiones/{id_ses}/vision', status_code=303)

    from .models import _query as _mq
    arch = _mq(
        "SELECT ArchivoPath AS archivo_path FROM RegistroDocumental WHERE IdRegistroDocumental=%s",
        (id_arch,), one=True
    )
    if not arch:
        return RedirectResponse(f'/cursos/{esp}/sesiones/{id_ses}/vision', status_code=303)

    ruta_archivo = os.path.join(BASE_DIR, 'static', arch['archivo_path'].replace('/', os.sep))
    try:
        from openai import OpenAI
        client = OpenAI(api_key=api_key)
        with open(ruta_archivo, 'rb') as audio_file:
            transcript = client.audio.transcriptions.create(
                model='whisper-1',
                file=audio_file,
                language='es',
            )
        texto = transcript.text
    except Exception as exc:
        request.session['flash_vision_ok'] = f'Error en transcripción: {exc}'
        return RedirectResponse(f'/cursos/{esp}/sesiones/{id_ses}/vision', status_code=303)

    sesion = get_sesion_by_id(id_ses)
    obs_actual = sesion.get('observaciones') or ''
    nueva_obs  = (obs_actual + '\n\n[Transcripción automática]\n' + texto).strip() if obs_actual else '[Transcripción automática]\n' + texto
    update_sesion_curso(id_ses, str(sesion['fecha']), sesion.get('resultado', ''),
                        nueva_obs, sesion.get('instructor', ''))
    request.session['flash_vision_ok'] = 'Transcripción guardada en observaciones de la sesión.'
    return RedirectResponse(f'/cursos/{esp}/sesiones/{id_ses}/vision', status_code=303)


from fastapi.responses import FileResponse as _FileResponse


@app.get('/vision/{id_arch}/exportar-word')
async def exportar_word_archivo(request: Request, id_arch: int):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    from .models import _query as _mq
    arch = _mq("""
        SELECT rd.ArchivoPath AS archivo_path, rd.TipoArchivo AS tipo, rd.Transcripcion AS notas,
               s.Fecha AS fecha, s.Especialidad AS especialidad, s.Resultado AS resultado,
               c.Nombre AS nombre_canino, c.Raza AS raza,
               e.RazonSocial AS nombre_entidad,
               pk.Nombre AS nombre_programa
        FROM RegistroDocumental rd
        JOIN Sesion s ON rd.IdSesion = s.IdSesion
        JOIN Canino c ON s.IdCanino = c.IdCanino
        LEFT JOIN Entidad e ON s.IdEntidad = e.IdEntidad
        LEFT JOIN ProgramaK9 pk ON s.Especialidad = pk.Clave
        WHERE rd.IdRegistroDocumental = %s
    """, (id_arch,), one=True)
    if not arch:
        return RedirectResponse('/cursos', status_code=303)

    from docx import Document
    from docx.shared import Pt, RGBColor, Inches, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    import tempfile

    GOLD = RGBColor(0xB8, 0x86, 0x0B)
    DARK = RGBColor(0x1A, 0x25, 0x40)
    GRAY = RGBColor(0x5A, 0x6B, 0x84)

    def _set_color(run, color): run.font.color.rgb = color
    def _titulo_doc(doc, texto):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(texto)
        r.bold = True; r.font.size = Pt(20); _set_color(r, GOLD)
    def _subtitulo(doc, texto):
        p = doc.add_paragraph()
        r = p.add_run(texto.upper())
        r.bold = True; r.font.size = Pt(11); _set_color(r, DARK)
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
    def _linea(doc):
        p = doc.add_paragraph('─' * 72)
        p.runs[0].font.color.rgb = RGBColor(0xDC, 0xE3, 0xEF)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)

    doc = Document()
    # Márgenes
    for sec in doc.sections:
        sec.top_margin = Cm(2); sec.bottom_margin = Cm(2)
        sec.left_margin = Cm(2.5); sec.right_margin = Cm(2.5)

    _titulo_doc(doc, 'CANGUARD')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Sistema de Gestión Canina K9 — Informe de Visión Artificial')
    r.font.size = Pt(10); _set_color(r, GRAY)
    _linea(doc)

    _subtitulo(doc, 'Datos del archivo')
    t = doc.add_table(rows=0, cols=2)
    t.style = 'Table Grid'
    res_map = {'aprobado': 'Aprobado', 'reprobado': 'Reprobado', 'en_proceso': 'En proceso'}
    for k, v in [
        ('Programa',  arch.get('nombre_programa') or arch.get('especialidad', '') or '—'),
        ('Canino',    arch.get('nombre_canino', '') or '—'),
        ('Raza',      arch.get('raza', '') or '—'),
        ('Entidad',   arch.get('nombre_entidad', '') or '—'),
        ('Fecha',     str(arch.get('fecha', '')) or '—'),
        ('Tipo',      (arch.get('tipo') or 'archivo').upper()),
        ('Resultado', res_map.get(arch.get('resultado', ''), arch.get('resultado', '') or '—')),
        ('Archivo',   arch.get('archivo_path', '').split('/')[-1]),
    ]:
        row = t.add_row()
        rc = row.cells[0].paragraphs[0].add_run(k)
        rc.bold = True; _set_color(rc, DARK)
        row.cells[1].text = str(v)
    for col_idx, width in [(0, Cm(4)), (1, Cm(12))]:
        for cell in t.columns[col_idx].cells:
            cell.width = width

    notas = (arch.get('notas') or '').strip()
    _subtitulo(doc, 'Transcripción completa')
    _linea(doc)
    if notas:
        for bloque in notas.replace('\r\n', '\n').split('\n'):
            bloque = bloque.strip()
            p = doc.add_paragraph(bloque if bloque else ' ')
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.first_line_indent = Cm(0.5)
    else:
        p = doc.add_paragraph('(Sin transcripción registrada)')
        _set_color(p.runs[0], GRAY); p.runs[0].italic = True

    _linea(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run('Generado por CANGUARD — Sistema K9')
    r.font.size = Pt(8); _set_color(r, GRAY)

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix='.docx')
    doc.save(tmp.name)
    canino = arch.get('nombre_canino', 'canino').lower().replace(' ', '_')
    nombre = f'transcripcion_{canino}_{id_arch}.docx'
    return _FileResponse(tmp.name,
        media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        filename=nombre)


@app.get('/cursos/{esp}/sesiones/{id_ses}/exportar-word')
async def exportar_word(request: Request, esp: str, id_ses: int):
    r = redirigir_si_no_autenticado(request)
    if r:
        return r
    sesion   = get_sesion_by_id(id_ses)
    archivos = get_vision_archivos(id_ses)
    curso    = get_programa_by_clave(esp)
    if not sesion:
        return RedirectResponse(f'/cursos/{esp}/sesiones', status_code=303)

    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    import tempfile

    GOLD = RGBColor(0xB8, 0x86, 0x0B)
    DARK = RGBColor(0x1A, 0x25, 0x40)
    GRAY = RGBColor(0x5A, 0x6B, 0x84)
    res_map = {'aprobado': 'Aprobado', 'reprobado': 'Reprobado', 'en_proceso': 'En proceso'}

    def _sc(run, color): run.font.color.rgb = color
    def _linea_doc(doc):
        p = doc.add_paragraph('─' * 72)
        p.runs[0].font.color.rgb = RGBColor(0xDC, 0xE3, 0xEF)
        p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(6)
    def _seccion(doc, texto):
        p = doc.add_paragraph()
        r = p.add_run(texto.upper())
        r.bold = True; r.font.size = Pt(11); _sc(r, DARK)
        p.paragraph_format.space_before = Pt(16); p.paragraph_format.space_after = Pt(4)

    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(2); sec.bottom_margin = Cm(2)
        sec.left_margin = Cm(2.5); sec.right_margin = Cm(2.5)

    # Portada
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('CANGUARD')
    r.bold = True; r.font.size = Pt(22); _sc(r, GOLD)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Informe de Sesión K9 — Visión Artificial')
    r.font.size = Pt(11); _sc(r, GRAY)
    _linea_doc(doc)

    # Datos de sesión
    _seccion(doc, 'Datos de la sesión')
    tabla = doc.add_table(rows=0, cols=2)
    tabla.style = 'Table Grid'
    for k, v in [
        ('Programa',     curso['nombre'] if curso else esp),
        ('Canino',       sesion.get('nombre_canino', '') or '—'),
        ('Raza',         sesion.get('raza', '') or '—'),
        ('Entidad',      sesion.get('nombre_entidad', '') or '—'),
        ('Fecha',        str(sesion.get('fecha', '')) or '—'),
        ('Guía Canino',  sesion.get('instructor', '') or '—'),
        ('Resultado',    res_map.get(sesion.get('resultado', ''), sesion.get('resultado', '') or 'Sin registro')),
        ('Archivos',     str(len(archivos))),
    ]:
        row = tabla.add_row()
        rc = row.cells[0].paragraphs[0].add_run(k)
        rc.bold = True; _sc(rc, DARK)
        row.cells[1].text = str(v)
    for cell in tabla.columns[0].cells: cell.width = Cm(4)
    for cell in tabla.columns[1].cells: cell.width = Cm(12)

    obs = (sesion.get('observaciones') or '').strip()
    if obs:
        _seccion(doc, 'Observaciones de la sesión')
        _linea_doc(doc)
        doc.add_paragraph(obs)

    # Transcripciones
    if archivos:
        _seccion(doc, f'Transcripciones ({len(archivos)} archivo{"s" if len(archivos) != 1 else ""})')
        _linea_doc(doc)
        for i, arch in enumerate(archivos, 1):
            # Encabezado de archivo
            p = doc.add_paragraph()
            r = p.add_run(f'Archivo {i}  —  {(arch.get("tipo") or "archivo").upper()}')
            r.bold = True; r.font.size = Pt(11); _sc(r, DARK)
            p.paragraph_format.space_before = Pt(12 if i > 1 else 0)

            fname = (arch.get('archivo_path') or '').split('/')[-1]
            p2 = doc.add_paragraph()
            r2 = p2.add_run(fname)
            r2.font.size = Pt(9); _sc(r2, GRAY); r2.italic = True
            p2.paragraph_format.space_after = Pt(6)

            notas = (arch.get('notas') or '').strip()
            if notas:
                for bloque in notas.replace('\r\n', '\n').split('\n'):
                    bloque = bloque.strip()
                    p = doc.add_paragraph(bloque if bloque else ' ')
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.first_line_indent = Cm(0.5)
            else:
                p = doc.add_paragraph('(Sin transcripción para este archivo)')
                _sc(p.runs[0], GRAY); p.runs[0].italic = True

            if i < len(archivos):
                doc.add_paragraph()

    _linea_doc(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run('Generado por CANGUARD — Sistema de Gestión Canina K9')
    r.font.size = Pt(8); _sc(r, GRAY)

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix='.docx')
    doc.save(tmp.name)
    canino = sesion.get('nombre_canino', 'canino').lower().replace(' ', '_')
    nombre_archivo = f'informe_{esp}_{id_ses}_{canino}.docx'
    return _FileResponse(tmp.name,
        media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        filename=nombre_archivo)
